from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "Guia_Pruebas_Video_RF_Shield_Planner.docx"
SCREENSHOTS = ROOT / "output" / "screenshots"

NAVY = "17324D"
BLUE = "2F6B9A"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
MID_GRAY = "D9D9D9"
TEXT = "20262E"
FIGURE_NUMBER = 0


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_font(run, name="Aptos", size=11, bold=False, color=TEXT, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    r = p.add_run(text)
    set_font(r)
    return p


def add_step(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    r = p.add_run(f"{number}.  {text}")
    set_font(r)
    return p


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, size=10, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[idx])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_font(r, size=9.5)
            if widths:
                cell.width = Inches(widths[idx])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_script(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=10.5, italic=True, color="34495E")
    return p


def add_figure(doc, filename, caption, width=6.35):
    global FIGURE_NUMBER
    path = SCREENSHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    FIGURE_NUMBER += 1
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(5)
    picture.paragraph_format.space_after = Pt(3)
    picture.paragraph_format.keep_with_next = True
    picture.add_run().add_picture(str(path), width=Inches(width))
    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(9)
    run = label.add_run(f"Figura {FIGURE_NUMBER}. {caption}")
    set_font(run, size=9, italic=True, color="52616B")
    return picture, label


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=9, color="66727D")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(TEXT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(29)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_p_pr = title_style._element.get_or_add_pPr()
title_borders = title_p_pr.find(qn("w:pBdr"))
if title_borders is not None:
    title_p_pr.remove(title_borders)

for name, size, before, after in (("Heading 1", 18, 14, 7), ("Heading 2", 14, 11, 5), ("Heading 3", 12, 8, 4)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer.paragraphs[0].text = "RF Shield Planner  |  Guía de demostración"
set_font(footer.paragraphs[0].runs[0], size=9, color="66727D")
add_page_number(footer.add_paragraph())

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(70)
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("SIMULADOR DE INGENIERÍA")
set_font(r, size=11, bold=True, color=BLUE)

p = doc.add_paragraph(style="Title")
p.add_run("Guía de pruebas y guion para el video de RF Shield Planner")
p.paragraph_format.space_after = Pt(18)
remove_paragraph_borders(p)

p = doc.add_paragraph()
r = p.add_run("Demostración de atenuación pasiva, balance de potencias y diseño de escenarios")
set_font(r, size=15, color="34495E")
p.paragraph_format.space_after = Pt(28)

add_body(doc, "Este documento indica exactamente qué configurar, qué botones usar, qué resultados observar y qué explicar durante un video de hasta cinco minutos. La demostración principal compara 850 MHz con 28 GHz y muestra por qué las frecuencias bajas requieren un cerramiento arquitectónico más exigente.")

add_body(doc, "Proyecto: RF Shield Planner", bold_lead="Proyecto:")
add_body(doc, "Tipo de entrega: guía complementaria para el video demostrativo", bold_lead="Tipo de entrega:")
add_body(doc, "Duración recomendada del video: 4 minutos 30 segundos", bold_lead="Duración recomendada del video:")

doc.add_page_break()

add_heading(doc, "Objetivo de la demostración", 1)
add_body(doc, "El video debe comprobar que el simulador implementa un balance de potencias coherente y que el plano modifica la potencia recibida según la distancia, la frecuencia, el material, el grosor y la cantidad de obstáculos atravesados. El resultado más importante es la comparación del mismo edificio a 850 MHz y a 28 GHz.")
add_body(doc, "La explicación debe mantener el alcance académico del proyecto. Los coeficientes de los materiales son didácticos y sirven para comparar escenarios. No representan por sí solos una especificación constructiva ni sustituyen mediciones de campo.")

add_heading(doc, "Ruta recomendada para un video menor de cinco minutos", 1)
add_table(doc,
          ["Tiempo", "Contenido en pantalla", "Idea que debe quedar clara"],
          [
              ["0:00 a 0:25", "Pantalla principal y estado API conectada", "El sistema separa frontend y backend y analiza atenuación pasiva"],
              ["0:25 a 1:05", "Código de las fórmulas y panel Balance calculado", "La potencia recibida resta FSPL y pérdidas de obstáculos"],
              ["1:05 a 2:05", "Ejemplo y comparación 850 MHz contra 28 GHz", "La misma arquitectura controla mejor 28 GHz"],
              ["2:05 a 3:25", "Escenario de 1900 MHz antes y después de agregar blindaje", "El material, grosor y continuidad cambian el confinamiento"],
              ["3:25 a 4:05", "Puerta, reja, selección, guardado o cambio de distancia", "El plano es editable y permite explorar escenarios"],
              ["4:05 a 4:35", "Resultados y cierre ético", "El proyecto no emite interferencia y requiere validación real"],
          ], widths=[1.05, 2.6, 2.85])

add_heading(doc, "Preparación antes de grabar", 1)
for item in [
    "Encender el backend en el puerto 3000 y el frontend en el puerto 5173.",
    "Confirmar que la esquina superior muestre API conectada.",
    "Abrir el simulador con zoom del navegador entre 90 y 100 por ciento.",
    "Tener abierto backend/src/rf-model.js para enseñar las funciones de cálculo.",
    "Pulsar Ejemplo antes de comenzar las pruebas numéricas.",
    "Desactivar notificaciones y comprobar el micrófono.",
    "Ensayar una vez con cronómetro. La grabación final no debe superar cinco minutos.",
]:
    add_bullet(doc, item)

doc.add_page_break()

add_heading(doc, "Prueba 1 Conexión y carga del escenario", 1)
add_body(doc, "Propósito: demostrar que el frontend obtiene los cálculos desde un backend independiente y que el editor carga un escenario reproducible.", bold_lead="Propósito:")
add_heading(doc, "Pasos en pantalla", 2)
for number, item in enumerate([
    "Mostrar el indicador API conectada.",
    "Pulsar Ejemplo para cargar el edificio, los muros, la puerta, la reja y la antena.",
    "Señalar que cada cuadro grande representa dos metros.",
    "Pulsar Simular cobertura y esperar el mapa de colores.",
], start=1):
    add_step(doc, number, item)
add_heading(doc, "Qué decir", 2)
add_script(doc, "El proyecto tiene dos aplicaciones separadas. El frontend contiene la interfaz y el editor del plano, mientras que el backend recibe el escenario y ejecuta el modelo de propagación. El indicador verde confirma que ambos componentes se están comunicando. Para que la demostración sea repetible cargo un plano de ejemplo, pero el usuario también puede construir su propio escenario.")
add_heading(doc, "Qué observar", 2)
add_bullet(doc, "El panel Resultado debe mostrar porcentaje de área bajo el umbral, potencia promedio, longitud de onda y potencia mínima.")
add_bullet(doc, "El panel Balance calculado debe mostrar distancia promedio, FSPL promedio y pérdida promedio por obstáculos.")
add_figure(
    doc,
    "00_pantalla_inicial.png",
    "Pantalla inicial del simulador con el frontend conectado al backend. En el video señale el indicador verde API conectada y los controles de frecuencia, potencia, distancia y umbral.",
    width=4.5,
)

add_heading(doc, "Prueba 2 Fórmula de atenuación y balance de potencias", 1)
add_body(doc, "Propósito: explicar la parte técnica exigida en el entregable sin convertir el video en una revisión completa del código.", bold_lead="Propósito:")
add_heading(doc, "Fórmulas implementadas", 2)
add_body(doc, "FSPL dB = 32.44 + 20 log10 distancia en km + 20 log10 frecuencia en MHz")
add_body(doc, "Pérdida del material dB = coeficiente de referencia × grosor × frecuencia en GHz elevada al exponente del material")
add_body(doc, "Potencia recibida dBm = potencia transmitida + ganancias - FSPL - suma de pérdidas por obstáculos")
add_heading(doc, "Qué mostrar en el código", 2)
for item in [
    "calculateFspl para la pérdida en espacio libre.",
    "calculateObstacleLoss para la pérdida del material según grosor y frecuencia.",
    "calculatePoint para la potencia recibida y la comparación con el umbral.",
]:
    add_bullet(doc, item)
add_heading(doc, "Qué decir", 2)
add_script(doc, "Primero calculo la pérdida en espacio libre. Esta pérdida aumenta con la distancia y con la frecuencia. Después identifico los elementos que cruza el trayecto entre la antena y cada punto del plano. Cada obstáculo aporta una pérdida según su material, grosor y dependencia con la frecuencia. Finalmente resto esas pérdidas a la potencia transmitida y comparo la potencia recibida con el umbral de menos 95 dBm. Si el resultado queda por debajo del umbral, el punto se considera confinado dentro de este modelo conceptual.")

add_heading(doc, "Prueba 3 Comparación principal entre 850 MHz y 28 GHz", 1)
add_body(doc, "Esta es la prueba central del análisis So What y conviene incluirla siempre en el video.")
add_heading(doc, "Configuración", 2)
add_table(doc,
          ["Parámetro", "Valor"],
          [
              ["Plano", "Ejemplo incluido"],
              ["Potencia transmitida", "43 dBm"],
              ["Distancia a la antena", "1.0 km"],
              ["Umbral de comunicación", "-95 dBm"],
              ["Acción", "Pulsar Comparar 850 MHz vs 28 GHz"],
          ], widths=[2.1, 4.4])
add_heading(doc, "Resultados de referencia", 2)
add_table(doc,
          ["Banda", "Área bajo el umbral", "Potencia promedio", "Longitud de onda"],
          [
              ["850 MHz", "0.0 por ciento", "Aproximadamente -56.7 dBm", "35.3 cm"],
              ["28 GHz", "100.0 por ciento", "Aproximadamente -124.7 dBm", "10.7 mm"],
          ], widths=[1.1, 1.6, 2.2, 1.6])
add_body(doc, "Los valores pueden variar ligeramente si el plano fue modificado. Lo importante es conservar el mismo plano y los mismos parámetros para ambas frecuencias.")
add_heading(doc, "Qué decir mientras se muestra 850 MHz", 2)
add_script(doc, "En 850 MHz la longitud de onda es de aproximadamente 35 centímetros. La potencia promedio se mantiene muy por encima del umbral de menos 95 dBm, por lo que el plano de ejemplo no logra confinar la señal. El resultado representa la mayor capacidad de penetración de una frecuencia baja y muestra que su control pasivo exige muros más gruesos, materiales con mayor pérdida o varias capas continuas.")
add_heading(doc, "Qué decir mientras se muestra 28 GHz", 2)
add_script(doc, "Con el mismo edificio y sin cambiar la potencia ni la distancia, 28 GHz llega mucho más débil. La longitud de onda se reduce a cerca de un centímetro, aumenta la pérdida en espacio libre y también aumenta la pérdida efectiva en los materiales del modelo. Por eso prácticamente toda el área queda por debajo del umbral. Esta comparación demuestra que las ondas milimétricas son más fáciles de confinar mediante barreras pasivas.")
add_heading(doc, "Conclusión breve de la prueba", 2)
add_script(doc, "La arquitectura no produce el mismo efecto en todas las bandas. Un diseño suficiente para 28 GHz puede ser insuficiente para 850 MHz. Por eso la frecuencia es una variable de diseño y no solamente un dato informativo.")

add_heading(doc, "Capturas de la prueba 850 MHz contra 28 GHz", 1)
add_body(doc, "Estas imágenes pueden mostrarse como evidencia durante la narración. Use el cursor para señalar primero el plano, después los indicadores numéricos y finalmente la diferencia entre bandas.")
add_figure(
    doc,
    "01_prueba_850_mhz_plano.png",
    "Plano y mapa de cobertura a 850 MHz. Señale las zonas cálidas y explique que el cerramiento del ejemplo todavía deja la potencia por encima de −95 dBm.",
    width=5.0,
)
doc.add_page_break()
add_figure(
    doc,
    "02_prueba_850_mhz_resultados.png",
    "Resultados a 850 MHz: 0.0 % del área bajo el umbral, potencia promedio de −56.7 dBm y longitud de onda de 35.3 cm.",
)
doc.add_page_break()
add_figure(
    doc,
    "03_comparacion_850_vs_28.png",
    "Comparación automática del mismo escenario: 850 MHz conserva 0.0 % bajo el umbral y 28 GHz alcanza 100.0 %. La diferencia promedio observada es de 68.0 dB.",
)
doc.add_page_break()
add_figure(
    doc,
    "04_prueba_28_ghz_mapa.png",
    "Mapa a 28 GHz. El predominio de colores oscuros representa potencias recibidas inferiores al umbral de comunicación.",
)

add_heading(doc, "Prueba 4 Material y grosor a 1900 MHz", 1)
add_body(doc, "Propósito: demostrar que la pérdida por material se acumula cuando el trayecto cruza una barrera y que el resultado cambia al modificar el material o el grosor.", bold_lead="Propósito:")
add_heading(doc, "Escenario base", 2)
for number, item in enumerate([
    "Pulsar Ejemplo.",
    "Seleccionar 3G 1900 MHz.",
    "Mantener 43 dBm, 1.0 km y -95 dBm.",
    "Pulsar Simular cobertura y anotar la potencia promedio y el área bajo el umbral.",
], start=1):
    add_step(doc, number, item)
add_heading(doc, "Escenario modificado", 2)
for number, item in enumerate([
    "Elegir Blindaje metálico y grosor de 0.25 m.",
    "Seleccionar Muro.",
    "Dibujar una barrera vertical continua dentro del cerramiento, cercana al muro izquierdo.",
    "Simular nuevamente y comparar el panel Balance calculado.",
], start=1):
    add_step(doc, number, item)
add_heading(doc, "Resultado de referencia", 2)
add_body(doc, "En la prueba capturada, una barrera vertical de blindaje metálico de aproximadamente 33.0 m y 0.25 m de grosor produjo 83.3 por ciento de área bajo el umbral y una potencia promedio de -108.7 dBm. La cifra exacta cambia con la posición, orientación y longitud del segmento dibujado.")
add_heading(doc, "Qué decir", 2)
add_script(doc, "La primera simulación funciona como línea base. Ahora agrego una barrera de blindaje metálico con 25 centímetros de grosor. El modelo encuentra qué trayectos atraviesan esa barrera y suma su pérdida al balance. Al repetir la simulación aumenta la pérdida promedio por obstáculos y disminuye la potencia recibida. Este valor de grosor es didáctico: sirve para hacer visible el comportamiento de la fórmula, no para recomendar una construcción real.")

add_heading(doc, "Capturas de la prueba con blindaje a 1900 MHz", 1)
add_figure(
    doc,
    "05_prueba_1900_base.png",
    "Escenario base a 1900 MHz antes de agregar la barrera adicional. Esta captura funciona como referencia para el antes y después.",
    width=5.0,
)
doc.add_page_break()
add_figure(
    doc,
    "06_blindaje_1900_mhz_mapa.png",
    "Barrera vertical de blindaje metálico dibujada dentro del cerramiento. Durante el video, siga con el cursor los trayectos que ahora cruzan el nuevo muro.",
)
doc.add_page_break()
add_figure(
    doc,
    "07_blindaje_1900_mhz_resultados.png",
    "Resultado capturado después del blindaje: 83.3 % bajo el umbral, −108.7 dBm de potencia promedio y 53.3 dB de pérdida promedio por obstáculos.",
)

add_heading(doc, "Prueba 5 Efecto de la distancia", 1)
add_body(doc, "Propósito: comprobar que FSPL cambia aunque el plano y los materiales permanezcan iguales.", bold_lead="Propósito:")
add_table(doc,
          ["Ejecución", "Distancia", "Qué debe ocurrir"],
          [
              ["A", "0.5 km", "Mayor potencia recibida y menor FSPL"],
              ["B", "2.0 km", "Menor potencia recibida y mayor FSPL"],
          ], widths=[1.0, 1.3, 4.2])
add_heading(doc, "Qué decir", 2)
add_script(doc, "Mantengo el mismo plano, la misma frecuencia y la misma potencia transmitida. Solo aumento la distancia a la antena. Como FSPL depende del logaritmo de la distancia, el panel muestra una pérdida mayor y la potencia recibida disminuye. Esta prueba permite separar el efecto de la propagación exterior del efecto de los materiales del edificio.")

add_heading(doc, "Prueba 6 Continuidad del cerramiento con puerta y reja", 1)
add_body(doc, "Propósito: mostrar que un acceso puede convertirse en el trayecto de menor atenuación y que el diseño debe analizar el cerramiento completo.", bold_lead="Propósito:")
add_heading(doc, "Pasos sugeridos", 2)
for number, item in enumerate([
    "Cargar el ejemplo y simular para conservar un resultado base.",
    "Usar Seleccionar y hacer clic en una puerta o reja para mostrar su tipo, longitud y grosor.",
    "Eliminar el elemento seleccionado y simular de nuevo para representar una abertura.",
    "Deshacer, volver a colocar la puerta o reja sobre el muro y repetir la simulación.",
], start=1):
    add_step(doc, number, item)
add_heading(doc, "Qué decir", 2)
add_script(doc, "Una barrera continua no depende únicamente del material de los muros. Las puertas, rejas y aberturas pueden dejar un trayecto con menos pérdida. El editor sustituye el tramo del muro cuando coloco una puerta o una reja, evitando contar dos materiales en el mismo lugar. Esta prueba ayuda a explicar por qué el confinamiento debe evaluarse como un sistema completo.")

add_heading(doc, "Dibujos para apoyar la explicación del video", 1)
add_body(doc, "Puede incluir estos dibujos como diapositivas breves o reproducirlos a mano. No necesitan animación: basta con mostrar cada uno durante diez o quince segundos y señalar las flechas mientras explica la idea.")
add_heading(doc, "Dibujo A Frecuencia y penetración", 2)
add_script(doc, "Aquí comparo el mismo muro. La onda de 850 MHz tiene una longitud de onda mayor y conserva más energía después del obstáculo. En 28 GHz la propagación se debilita con mayor facilidad al encontrar una superficie sólida. Por eso el mismo edificio no ofrece el mismo confinamiento en ambas bandas.")
add_figure(
    doc,
    "dibujo_01_frecuencia_penetracion.png",
    "Esquema conceptual para explicar por qué 850 MHz atraviesa obstáculos con mayor facilidad que 28 GHz.",
    width=5.6,
)
doc.add_page_break()
add_heading(doc, "Dibujo B Cerramiento continuo y punto débil", 2)
add_script(doc, "El material del muro no es suficiente si el cerramiento tiene una abertura. La señal busca el trayecto con menor pérdida. Por eso las uniones, puertas y rejas deben analizarse junto con todos los muros y no como piezas aisladas.")
add_figure(
    doc,
    "dibujo_02_continuidad_cerramiento.png",
    "Comparación entre una envolvente continua y otra con una abertura que funciona como punto de fuga.",
)
doc.add_page_break()
add_heading(doc, "Dibujo C Cadena del balance de potencias", 2)
add_script(doc, "Parto de la potencia transmitida. Luego resto la pérdida en espacio libre y las pérdidas acumuladas de los obstáculos. El resultado es la potencia recibida. Si queda por debajo de menos 95 dBm, el simulador clasifica ese punto como área bajo el umbral de comunicación.")
add_figure(
    doc,
    "dibujo_03_balance_potencias.png",
    "Secuencia visual del cálculo: potencia transmitida, FSPL, obstáculos, potencia recibida y comparación con el umbral.",
)

add_heading(doc, "Prueba 7 Edición y persistencia del plano", 1)
add_body(doc, "Esta prueba es opcional si queda tiempo. Sirve para enseñar que el simulador también funciona como herramienta de diseño conceptual.")
for item in [
    "Dibujar un muro y comprobar que aparece su longitud.",
    "Seleccionarlo y cambiar su material o grosor.",
    "Mover un elemento para reorganizar el escenario.",
    "Pulsar Guardar, luego Limpiar y finalmente Cargar.",
]:
    add_bullet(doc, item)
add_script(doc, "Además de ejecutar un cálculo fijo, la aplicación permite construir un plano conceptual. Puedo agregar muros, puertas, rejas, zonas interiores y cambiar la posición de la antena. El plano se puede guardar en el navegador y recuperar después para continuar una prueba.")

add_heading(doc, "Tabla de coeficientes didácticos del simulador", 1)
add_table(doc,
          ["Material", "Coeficiente de referencia", "Exponente de frecuencia", "Uso principal"],
          [
              ["Concreto reforzado", "22 dB por metro", "0.50", "Muros"],
              ["Ladrillo", "11 dB por metro", "0.65", "Muros"],
              ["Blindaje metálico", "145 dB por metro", "0.48", "Refuerzo de muros"],
              ["Puerta metálica", "48 dB por metro", "0.38", "Accesos"],
              ["Malla de acero", "34 dB por metro", "0.42", "Rejas"],
          ], widths=[1.6, 1.7, 1.5, 1.7])
add_body(doc, "Estos coeficientes están centralizados en el backend. Deben citarse como valores didácticos del prototipo y validarse con bibliografía técnica o mediciones antes de cualquier uso de ingeniería real.")

add_heading(doc, "Guion completo listo para narrar", 1)
add_heading(doc, "Inicio de 0:00 a 0:25", 2)
add_script(doc, "Este proyecto presenta RF Shield Planner, un simulador conceptual de atenuación pasiva para centros penitenciarios. La propuesta no utiliza inhibidores ni transmite interferencia. Su objetivo es estudiar cómo la frecuencia, la distancia y los materiales constructivos afectan la potencia que llega al interior de un edificio.")

add_heading(doc, "Arquitectura y fórmula de 0:25 a 1:05", 2)
add_script(doc, "La aplicación está separada en frontend y backend. El frontend contiene el editor del plano y presenta los resultados. El backend ejecuta el modelo. Primero calcula FSPL con la distancia en kilómetros y la frecuencia en megahercios. Después calcula la pérdida de cada obstáculo con un coeficiente por metro, el grosor y un factor de frecuencia. La potencia recibida se obtiene restando FSPL y las pérdidas de obstáculos a la potencia transmitida. Finalmente se compara con el umbral de menos 95 dBm.")

add_heading(doc, "Comparación de 1:05 a 2:05", 2)
add_script(doc, "Cargo el plano de ejemplo, mantengo 43 dBm y una distancia de un kilómetro, y ejecuto la comparación automática. En 850 MHz la longitud de onda es cercana a 35 centímetros y la potencia promedio permanece por encima del umbral. El cerramiento del ejemplo no confina la señal. En 28 GHz, con exactamente el mismo plano, la potencia llega mucho más débil y prácticamente toda el área queda bajo el umbral. El aumento de FSPL y de la pérdida efectiva en los materiales explica por qué las ondas milimétricas son más fáciles de controlar mediante blindaje pasivo.")

add_heading(doc, "Escenario construido de 2:05 a 3:25", 2)
add_script(doc, "Ahora selecciono 1900 MHz y ejecuto una línea base. Después agrego una barrera continua de blindaje metálico con 25 centímetros de grosor. Al simular otra vez, el balance muestra una mayor pérdida por obstáculos y una menor potencia promedio. El área bajo el umbral aumenta porque los trayectos que cruzan la nueva barrera acumulan más atenuación. El grosor utilizado es didáctico y permite demostrar la fórmula; no constituye una recomendación constructiva.")

add_heading(doc, "Editor de 3:25 a 4:05", 2)
add_script(doc, "El plano también permite representar muros, puertas, rejas, pisos y la antena. Al colocar una puerta o una reja sobre un muro, el editor sustituye ese tramo para no sumar dos materiales de forma incorrecta. También puedo seleccionar elementos, cambiar su grosor, moverlos, eliminarlos y guardar el escenario en el navegador.")

add_heading(doc, "Ética y cierre de 4:05 a 4:35", 2)
add_script(doc, "La solución propuesta se limita a la atenuación pasiva. Esto evita emitir energía interferente y protege las comunicaciones legítimas del exterior, incluidos los servicios de emergencia y los derechos de terceros que usan el espectro. El simulador demuestra una tendencia de diseño, pero una implementación real necesitaría mediciones, fuentes técnicas, análisis regulatorio y autorización institucional. En conclusión, 850 MHz exige una arquitectura más pesada, mientras que 28 GHz se confina con mayor facilidad debido a sus mayores pérdidas de propagación y penetración.")

add_heading(doc, "Preguntas que podrían hacer durante la evaluación", 1)
qa_rows = [
    ["¿Por qué usa -95 dBm?", "Es el umbral conceptual definido para clasificar si un punto conserva comunicación. Puede modificarse para estudiar otros criterios."],
    ["¿Por qué 28 GHz queda más bloqueada?", "Porque presenta mayor FSPL y, en el modelo, mayor pérdida efectiva al atravesar materiales. Su longitud de onda también es mucho menor."],
    ["¿El simulador diseña una cárcel real?", "No. Compara escenarios conceptuales. Un diseño real requiere planos detallados, mediciones, normas, bibliografía y validación profesional."],
    ["¿Los coeficientes son valores certificados?", "No. Son valores didácticos centralizados en el backend. El informe debe indicar que necesitan validación con fuentes o ensayos."],
    ["¿Por qué no usar un inhibidor?", "Porque un inhibidor transmite interferencia y puede afectar comunicaciones legítimas. El proyecto estudia exclusivamente control pasivo."],
    ["¿Qué aporta separar frontend y backend?", "La interfaz puede cambiar o publicarse por separado, mientras que el modelo de cálculo queda centralizado, validado y reutilizable."],
    ["¿Qué representa el mapa de colores?", "La potencia recibida estimada en cada punto. Los colores permiten localizar áreas por encima o por debajo del umbral."],
]
add_table(doc, ["Pregunta", "Respuesta sugerida"], qa_rows, widths=[2.1, 4.4])

add_heading(doc, "Errores comunes durante la grabación", 1)
add_table(doc,
          ["Problema", "Cómo resolverlo"],
          [
              ["API desconectada", "Encender el backend en el puerto 3000 y recargar el navegador"],
              ["No aparece el mapa", "Cargar Ejemplo y pulsar Simular cobertura"],
              ["Los resultados no coinciden con la guía", "Restablecer el ejemplo y verificar 43 dBm, 1.0 km y -95 dBm"],
              ["La barrera no cambia mucho el resultado", "Dibujarla de forma continua y comprobar que cruza varios trayectos desde la antena"],
              ["El video pasa de cinco minutos", "Omitir las pruebas 5, 6 y 7 y conservar comparación, fórmula, blindaje y ética"],
          ], widths=[2.0, 4.5])

add_heading(doc, "Lista final de verificación", 1)
for item in [
    "La grabación muestra API conectada.",
    "Se explica FSPL y la pérdida por material.",
    "Se muestra el balance de potencias en pantalla.",
    "Se compara el mismo plano en 850 MHz y 28 GHz.",
    "Se demuestra al menos una modificación del plano o del material.",
    "Se aclara que los coeficientes son didácticos.",
    "Se incluye la justificación ética y el alcance conceptual.",
    "La duración total es menor de cinco minutos.",
]:
    add_bullet(doc, "☐ " + item)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
