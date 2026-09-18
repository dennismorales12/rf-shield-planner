from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "Guion_Demostracion_5_Min_RF_Shield_Planner.docx"
SCREENSHOT = Path(r"C:\Users\hunte\AppData\Local\Temp\codex-clipboard-616b52bb-1345-416d-b627-ae2902c0f192.png")

NAVY = "0B2B44"
CYAN = "2FB7E8"
PALE_BLUE = "EAF6FB"
PALE_GRAY = "F4F6F8"
MID_GRAY = "D9D9D9"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(82, 96, 109)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Aptos"):
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        r_fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().append(r_fonts)
    else:
        r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        set_run_font(lead)
        body = p.add_run(text[len(bold_lead):])
        set_run_font(body)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.72)
    p.paragraph_format.first_line_indent = Cm(-0.48)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"{number}.\t{text}")
    set_run_font(r)
    return p


def add_check(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.62)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("☐  " + text)
    set_run_font(r)
    return p


def add_script(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.right_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.italic = True
    r.font.color.rgb = TEXT
    set_run_font(r)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, (cell, label) in enumerate(zip(header.cells, headers)):
        set_cell_fill(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(font_size)
        set_run_font(r)
        if widths:
            cell.width = Inches(widths[idx])
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            set_cell_fill(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(value)
            r.font.size = Pt(font_size)
            r.font.color.rgb = TEXT
            set_run_font(r)
            if widths:
                cell.width = Inches(widths[idx])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def set_image_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.55)
section.bottom_margin = Cm(1.45)
section.left_margin = Cm(1.75)
section.right_margin = Cm(1.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.2)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style.font.size = Pt(23)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(5)
title_ppr = title_style.element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name, size, before, after in (
    ("Heading 1", 15, 12, 5),
    ("Heading 2", 11.5, 8, 3),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Number"):
    styles[list_style].font.name = "Aptos"
    styles[list_style].font.size = Pt(10)
    styles[list_style].font.color.rgb = TEXT
    styles[list_style].paragraph_format.space_after = Pt(3)

doc.core_properties.title = "Guion de demostración de RF Shield Planner en 5 minutos"
doc.core_properties.subject = "Explicación y ejemplos para una demostración breve del simulador"
doc.core_properties.keywords = "RF, telecomunicaciones, atenuación, FSPL, simulador, demostración"

# Página 1
p = doc.add_paragraph(style="Title")
p.add_run("Guion de demostración de RF Shield Planner en 5 minutos")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(9)
r = p.add_run("Objetivo")
r.bold = True
r.font.color.rgb = RGBColor.from_string(NAVY)
set_run_font(r)
r = p.add_run("  Explicar qué resuelve la aplicación, cómo interpreta la cobertura y demostrar dos cambios de escenario sin superar cinco minutos.")
set_run_font(r)

if SCREENSHOT.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(SCREENSHOT), width=Inches(6.75))
    set_image_alt_text(
        shape,
        "Interfaz de RF Shield Planner",
        "Editor de plano con controles de tecnología y materiales, mapa de potencia recibida y panel de resultados.",
    )
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    run = caption.add_run("Figura 1  Interfaz del simulador en un escenario de 1900 MHz")
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED
    set_run_font(run)

add_heading(doc, "Explicación en una frase", 1)
add_script(doc, "RF Shield Planner es un simulador didáctico que estima cuánta señal celular llega a cada punto de un edificio según la frecuencia, la distancia a la antena y las pérdidas causadas por muros, puertas, rejas y materiales.")

add_heading(doc, "Idea principal para defender", 1)
add_body(doc, "La aplicación permite comparar alternativas antes de una evaluación real. Su aporte es visualizar tendencias: una frecuencia baja suele penetrar mejor, mientras que una frecuencia alta y un cerramiento continuo suelen producir más atenuación. Los resultados son conceptuales y deben validarse con mediciones y fuentes técnicas antes de tomar decisiones de ingeniería.")

doc.add_page_break()

# Página 2
add_heading(doc, "Qué se puede explicar de la aplicación", 1)
add_table(
    doc,
    ["Parte", "Qué hace", "Qué conviene decir"],
    [
        ["Escenario RF", "Configura tecnología, potencia transmitida, distancia y umbral.", "Estas variables determinan la pérdida de propagación y el criterio para decidir si una zona conserva comunicación."],
        ["Editor del plano", "Permite dibujar muros, puertas, rejas, pisos o zonas y mover la antena.", "Cada elemento puede cambiar el recorrido y la pérdida acumulada de la señal."],
        ["Mapa de colores", "Muestra la potencia recibida estimada en cada cuadrícula.", "Rojo representa señal más fuerte; amarillo y azul muestran debilitamiento; gris oscuro indica puntos bajo el umbral."],
        ["Resultado", "Resume área bajo el umbral, potencia promedio, potencia mínima y longitud de onda.", "El promedio no describe todos los puntos: pueden existir zonas bloqueadas aunque la potencia promedio esté sobre el umbral."],
        ["Balance calculado", "Desglosa potencia transmitida, FSPL y pérdida por obstáculos.", "Este panel permite justificar el resultado y no limitarse a observar colores."],
    ],
    widths=[1.25, 2.25, 3.45],
    font_size=8.8,
)

add_heading(doc, "Cómo explicar el cálculo sin complicarlo", 1)
add_body(doc, "Potencia recibida = potencia transmitida + ganancias - pérdida en espacio libre - pérdidas por obstáculos.")
add_body(doc, "La pérdida en espacio libre o FSPL aumenta cuando crecen la distancia o la frecuencia. Después, el sistema identifica los muros, puertas o rejas que atraviesa el trayecto desde la antena y suma la pérdida correspondiente a su material y grosor.")

add_heading(doc, "Lectura del escenario mostrado", 1)
add_bullet(doc, "Frecuencia: 1900 MHz. La longitud de onda calculada es 15.8 cm.")
add_bullet(doc, "Potencia transmitida: 43 dBm. Distancia configurada: 2.0 km. Umbral: -95 dBm.")
add_bullet(doc, "El 37.2 % del área está bajo el umbral, por eso el confinamiento se clasifica como parcial.")
add_bullet(doc, "La potencia promedio es -84.8 dBm, pero la mínima llega a -132.3 dBm; el mapa contiene zonas con comportamientos distintos.")
add_bullet(doc, "El balance mostrado coincide: 43.0 - 104.3 - 23.5 = -84.8 dBm, con ganancias de antena configuradas en 0 dBi.")

add_script(doc, "Aquí no debo decir que todo el edificio está bloqueado. El promedio queda por encima de -95 dBm y solo una parte del área cae bajo el umbral. Esto ayuda a localizar accesos o segmentos que necesitan mayor atenuación.")

doc.add_page_break()

# Página 3
add_heading(doc, "Guion cronometrado listo para narrar", 1)
add_table(
    doc,
    ["Tiempo", "Acción en pantalla", "Narración sugerida"],
    [
        ["0:00 a 0:25", "Mostrar la pantalla completa y el indicador API conectada.", "Este es RF Shield Planner, un simulador conceptual de atenuación pasiva. Permite estudiar cómo la frecuencia, la distancia y los materiales de un edificio afectan la potencia de una señal celular."],
        ["0:25 a 1:00", "Señalar controles, herramientas del plano y panel Resultado.", "A la izquierda configuro el escenario. En el centro dibujo el edificio y observo el mapa de cobertura. A la derecha reviso los indicadores y el balance de potencia calculado por el backend."],
        ["1:00 a 1:40", "Explicar los colores y señalar los números de la captura.", "En este escenario de 1900 MHz, 37.2 por ciento del área está bajo -95 dBm. La potencia promedio es -84.8 dBm, pero el mínimo llega a -132.3 dBm. Por eso hay zonas útiles y zonas confinadas dentro del mismo plano."],
        ["1:40 a 2:40", "Pulsar Ejemplo y luego Comparar 850 MHz vs 28 GHz.", "Mantengo exactamente el mismo plano para aislar el efecto de la frecuencia. En 850 MHz la señal penetra mejor. En 28 GHz aumentan las pérdidas y una proporción mayor del área queda bajo el umbral."],
        ["2:40 a 3:35", "Cambiar solo la distancia de 0.5 km a 2.0 km y simular cada vez.", "Ahora mantengo la frecuencia, la potencia y el plano. Al aumentar cuatro veces la distancia, la FSPL aumenta aproximadamente 12 dB y la potencia recibida disminuye. Así separo el efecto de la distancia del efecto de los materiales."],
        ["3:35 a 4:10", "Señalar Balance calculado y una puerta o reja.", "El resultado se obtiene restando la pérdida en espacio libre y las pérdidas de los obstáculos. Una abertura o un acceso con menor atenuación puede convertirse en un punto débil, aunque los muros sean resistentes."],
        ["4:10 a 4:35", "Volver a la vista general y cerrar.", "La aplicación no reemplaza mediciones de campo ni entrega un diseño constructivo definitivo. Sirve para comparar escenarios y demostrar que el confinamiento depende de la frecuencia, la distancia, el material, el grosor y la continuidad del cerramiento."],
    ],
    widths=[0.95, 2.0, 4.0],
    font_size=8.35,
)
add_body(doc, "Duración prevista: 4 minutos 35 segundos. Quedan aproximadamente 25 segundos de margen para esperar el cálculo o corregir una selección.", bold_lead="Duración prevista:")

add_heading(doc, "Frases cortas para señalar la pantalla", 1)
add_bullet(doc, "El mapa no muestra paredes fuertes o débiles; muestra la potencia estimada que llega a cada punto.")
add_bullet(doc, "Área bajo el umbral significa porcentaje de la zona evaluada con potencia menor que el valor definido.")
add_bullet(doc, "La potencia mínima localiza el punto más desfavorable; la potencia promedio resume todo el interior.")
add_bullet(doc, "El mismo edificio puede comportarse de forma muy distinta según la banda de frecuencia.")

doc.add_page_break()

# Página 4
add_heading(doc, "Ejemplos para la demostración", 1)
add_heading(doc, "Ejemplo 1 Comparar 850 MHz y 28 GHz", 2)
add_body(doc, "Esta es la prueba principal porque cambia una sola variable y produce una diferencia fácil de observar.")
for number, text in enumerate((
    "Pulse Ejemplo para recuperar el plano estándar.",
    "Configure 43 dBm, 1.0 km y umbral de -95 dBm.",
    "Pulse Comparar 850 MHz vs 28 GHz.",
    "Señale el porcentaje bajo el umbral y la potencia promedio de cada banda.",
), start=1):
    add_number(doc, number, text)
add_table(
    doc,
    ["Banda", "Resultado de referencia", "Interpretación"],
    [
        ["850 MHz", "0.0 % bajo el umbral y cerca de -56.7 dBm de promedio.", "La banda baja conserva más señal dentro del ejemplo."],
        ["28 GHz", "100.0 % bajo el umbral y cerca de -124.7 dBm de promedio.", "La frecuencia alta presenta mucha mayor atenuación en el mismo escenario."],
    ],
    widths=[1.15, 2.65, 3.15],
    font_size=8.8,
)
add_body(doc, "Los valores de referencia corresponden al plano estándar. Si el plano se modifica, cambian los resultados; la conclusión válida surge de comparar ambas bandas con el mismo plano y los mismos parámetros.")
add_script(doc, "Esta comparación muestra por qué la frecuencia es una variable de diseño. Una solución que confina 28 GHz puede ser insuficiente para 850 MHz.")

add_heading(doc, "Ejemplo 2 Cambiar la distancia", 2)
for number, text in enumerate((
    "Mantenga un solo plano, una sola frecuencia y 43 dBm.",
    "Simule primero con 0.5 km y observe FSPL y potencia promedio.",
    "Cambie únicamente la distancia a 2.0 km y vuelva a simular.",
    "Explique que cuadruplicar la distancia agrega aproximadamente 12 dB de pérdida en espacio libre.",
), start=1):
    add_number(doc, number, text)
add_script(doc, "Como no cambié el edificio, la diferencia proviene de la propagación exterior. Una antena más lejana produce mayor FSPL y menor potencia recibida.")

add_heading(doc, "Ejemplo 3 Agregar una barrera", 2)
add_body(doc, "Use este ejemplo como alternativa si desea una demostración más visual; no es necesario ejecutarlo además de los dos anteriores.")
for number, text in enumerate((
    "Cargue Ejemplo y seleccione 1900 MHz.",
    "Ejecute una simulación base y observe Pérdida por obstáculos.",
    "Seleccione Blindaje metálico, grosor de 0.25 m y herramienta Muro.",
    "Dibuje una barrera vertical continua dentro del cerramiento y simule otra vez.",
), start=1):
    add_number(doc, number, text)
add_script(doc, "La barrera aumenta la pérdida de los trayectos que la atraviesan. El grosor de 25 centímetros se usa para visualizar el efecto del modelo y no constituye una recomendación constructiva real.")

doc.add_page_break()

# Página 5
add_heading(doc, "Preguntas probables y respuestas breves", 1)
add_table(
    doc,
    ["Pregunta", "Respuesta sugerida"],
    [
        ["¿Qué significa -95 dBm?", "Es el umbral conceptual usado para clasificar si un punto conserva comunicación. Puede modificarse para probar otros criterios."],
        ["¿Por qué 28 GHz se bloquea más?", "Porque su FSPL es mayor y el modelo también aumenta la pérdida efectiva de los materiales con la frecuencia."],
        ["¿Qué representa el porcentaje?", "La proporción de puntos evaluados dentro de las zonas dibujadas que queda por debajo del umbral."],
        ["¿Por qué el promedio puede ser mayor que -95 dBm si hay área bloqueada?", "Porque el promedio combina puntos fuertes y débiles. El porcentaje y la potencia mínima muestran la distribución espacial."],
        ["¿Los resultados sirven para construir?", "No directamente. Son resultados didácticos que requieren mediciones, bibliografía, normativa y validación profesional."],
        ["¿La app genera interferencia?", "No. Modela atenuación pasiva; no transmite ni controla equipos de radio."],
    ],
    widths=[2.25, 4.7],
    font_size=8.8,
)

add_heading(doc, "Qué no conviene afirmar", 1)
add_bullet(doc, "No diga que el simulador garantiza bloqueo total de llamadas.")
add_bullet(doc, "No presente los coeficientes de los materiales como valores certificados.")
add_bullet(doc, "No convierta un escenario didáctico en una recomendación de construcción.")
add_bullet(doc, "No confunda potencia promedio con cobertura uniforme en todo el plano.")

add_heading(doc, "Lista antes de grabar", 1)
for item in (
    "Backend y frontend encendidos; el indicador debe mostrar API conectada.",
    "Plano de ejemplo cargado y controles visibles sin desplazamiento lateral.",
    "Zoom del navegador entre 90 y 100 por ciento.",
    "Notificaciones y ventanas emergentes desactivadas.",
    "Pruebas ensayadas una vez y resultados principales anotados.",
    "Cronómetro preparado y duración objetivo de 4 minutos 35 segundos.",
):
    add_check(doc, item)

add_heading(doc, "Cierre recomendado", 1)
add_script(doc, "RF Shield Planner convierte un balance de radiofrecuencia en un mapa fácil de interpretar. Su valor está en comparar escenarios de manera controlada y mostrar que la cobertura cambia por la combinación de frecuencia, distancia y cerramiento. El siguiente paso para un proyecto real sería validar el modelo con mediciones y datos técnicos de materiales.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
