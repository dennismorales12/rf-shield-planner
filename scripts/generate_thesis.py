from pathlib import Path
import math

from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "docx" / "Proyecto_Tesis_RF_Shield_Planner.docx"
ASSETS = ROOT / "tmp" / "thesis_assets"
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

BLACK = "000000"
DARK_BLUE = "17365D"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F2F2"
BORDER_GRAY = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 14, 12, 12),
        ("Heading 2", 12, 10, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Table Caption" not in [style.name for style in doc.styles]:
        caption = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = doc.styles["Table Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.bold = True
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=10)


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.different_first_page_header_footer = True
    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_no_indent(doc, text, alignment=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=12):
    paragraph = doc.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def add_chapter(doc, number, title):
    doc.add_page_break()
    paragraph = doc.add_paragraph(style="Heading 1")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(f"CAPÍTULO {number}\n{title.upper()}")
    set_font(run, size=14, bold=True)
    return paragraph


def add_table(doc, headers, rows, widths=None, trailing_space=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, DARK_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(text))
        set_font(run, size=9.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = widths[index]
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index > 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_font(run, size=9)
            if widths:
                cell.width = widths[col_index]
    if trailing_space:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(1.25)
    paragraph.paragraph_format.first_line_indent = Cm(-0.63)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_font(run)
    return paragraph


def make_result_chart():
    path = ASSETS / "resultados_escenarios.png"
    labels = ["850 MHz", "1900 MHz", "28 GHz", "1900 MHz\ncon blindaje"]
    values = [0, 0, 100, 94.4]
    colors = ["#8C8C8C", "#8C8C8C", "#2F75B5", "#17365D"]
    image = Image.new("RGB", (1600, 820), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
    small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 26)
    bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 32)
    left, top, right, bottom = 180, 80, 1510, 650
    draw.text((800, 25), "Área interior bajo el umbral de -95 dBm", font=bold, fill="black", anchor="ma")
    for value in range(0, 101, 20):
        y = bottom - (value / 100) * (bottom - top)
        draw.line((left, y, right, y), fill="#D9D9D9", width=2)
        draw.text((left - 20, y), f"{value} %", font=small, fill="black", anchor="rm")
    draw.line((left, top, left, bottom), fill="black", width=3)
    draw.line((left, bottom, right, bottom), fill="black", width=3)
    slot = (right - left) / len(labels)
    bar_width = 190
    for index, (label, value, color) in enumerate(zip(labels, values, colors)):
        center = left + slot * (index + 0.5)
        height = (value / 100) * (bottom - top)
        draw.rectangle((center - bar_width / 2, bottom - height, center + bar_width / 2, bottom), fill=color, outline="black", width=2)
        draw.text((center, max(top + 10, bottom - height - 16)), f"{value:g} %", font=small, fill="black", anchor="ms")
        if "\n" in label:
            first, second = label.split("\n")
            draw.text((center, bottom + 25), first, font=regular, fill="black", anchor="ma")
            draw.text((center, bottom + 60), second, font=regular, fill="black", anchor="ma")
        else:
            draw.text((center, bottom + 25), label, font=regular, fill="black", anchor="ma")
    image.save(path, dpi=(220, 220))
    return path


def add_figure(doc, image_path, caption, width=Inches(6.0)):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=width)
    caption_paragraph = doc.add_paragraph(caption, style="Table Caption")
    return caption_paragraph


def add_equation(doc, expression, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_para_pr.append(justification)
    math_para.append(math_para_pr)
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = expression
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    paragraph._p.append(math_para)
    label = doc.add_paragraph(caption)
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.first_line_indent = Cm(0)
    label.paragraph_format.space_after = Pt(8)
    set_font(label.runs[0], size=10, italic=True)


def coefficient_rows():
    frequencies = [850, 1900, 2100, 3500, 28000]
    materials = [
        ("Concreto reforzado", 22, 0.50),
        ("Ladrillo", 11, 0.65),
        ("Malla de acero", 34, 0.42),
        ("Puerta metálica", 48, 0.38),
        ("Blindaje metálico", 145, 0.48),
    ]
    rows = []
    for name, reference, exponent in materials:
        values = [reference * math.pow(frequency / 1000, exponent) for frequency in frequencies]
        rows.append([name] + [f"{value:.1f}" for value in values])
    return rows


def build_document():
    chart = make_result_chart()

    doc = Document()
    configure_page(doc)
    configure_styles(doc)

    # Portada
    for _ in range(2):
        doc.add_paragraph()
    add_no_indent(doc, "UNIVERSIDAD ________________________________________", WD_ALIGN_PARAGRAPH.CENTER, True, 14)
    add_no_indent(doc, "FACULTAD __________________________________________", WD_ALIGN_PARAGRAPH.CENTER, True, 12)
    add_no_indent(doc, "CARRERA ___________________________________________", WD_ALIGN_PARAGRAPH.CENTER, True, 12)
    doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.add_run(
        "Diseño de un simulador web para el análisis conceptual de atenuación pasiva "
        "de señales celulares en centros penitenciarios"
    )
    doc.add_paragraph()
    add_no_indent(doc, "PROYECTO DE GRADUACIÓN", WD_ALIGN_PARAGRAPH.CENTER, True, 12)
    add_no_indent(doc, "Presentado por", WD_ALIGN_PARAGRAPH.CENTER, False, 12)
    add_no_indent(doc, "________________________________________", WD_ALIGN_PARAGRAPH.CENTER, True, 12)
    add_no_indent(doc, "Carné __________________________________", WD_ALIGN_PARAGRAPH.CENTER, False, 12)
    doc.add_paragraph()
    add_no_indent(doc, "Asesor __________________________________", WD_ALIGN_PARAGRAPH.CENTER, False, 12)
    doc.add_paragraph()
    add_no_indent(doc, "Guatemala, ______________________________", WD_ALIGN_PARAGRAPH.CENTER, False, 12)

    # Hoja de identificación
    doc.add_page_break()
    add_no_indent(doc, "Datos del proyecto", WD_ALIGN_PARAGRAPH.CENTER, True, 14)
    metadata = [
        ("Título", "Diseño de un simulador web para el análisis conceptual de atenuación pasiva de señales celulares en centros penitenciarios"),
        ("Autor", "________________________________________"),
        ("Carné", "________________________________________"),
        ("Curso", "Telecomunicaciones"),
        ("Docente o asesor", "________________________________________"),
        ("Línea de investigación", "Propagación radioeléctrica y protección pasiva del espectro"),
        ("Modalidad", "Proyecto tecnológico con simulación web"),
    ]
    add_table(doc, ["Campo", "Información"], metadata, [Inches(1.75), Inches(4.75)])
    add_body(
        doc,
        "El documento presenta una propuesta académica y conceptual. No describe la construcción ni el uso de "
        "inhibidores activos de señal. Cualquier aplicación física requerirá mediciones, autorización institucional "
        "y coordinación con la Superintendencia de Telecomunicaciones y los operadores involucrados.",
        "El documento presenta",
    )

    # Resumen y abstract
    doc.add_page_break()
    add_heading(doc, "Resumen", 1)
    add_body(
        doc,
        "La conectividad celular no autorizada dentro de centros penitenciarios representa un problema de seguridad "
        "que no puede abordarse responsablemente mediante interferencia indiscriminada. Esta investigación desarrolla "
        "RF Shield Planner, un simulador web que permite dibujar un plano bidimensional y estimar la potencia recibida "
        "después de aplicar pérdida de espacio libre y atenuación por materiales constructivos. El sistema separa una "
        "interfaz desarrollada con HTML, CSS, JavaScript y Canvas de un backend en Node.js y Express encargado de los "
        "cálculos radioeléctricos.",
    )
    add_body(
        doc,
        "La metodología utiliza escenarios comparables con potencia transmitida de 43 dBm, distancia base de un "
        "kilómetro y umbral conceptual de -95 dBm. Los resultados muestran que un cerramiento convencional mantiene "
        "la señal de 850 MHz por encima del umbral, mientras que el mismo plano produce confinamiento total a 28 GHz "
        "dentro de los parámetros del modelo. Un refuerzo metálico conceptual aplicado a 1900 MHz aumenta el área "
        "confinada de 0 % a aproximadamente 94.4 %. Se concluye que las frecuencias bajas requieren mayor masa, varias "
        "capas y control estricto de aberturas, mientras las ondas milimétricas presentan un presupuesto de enlace más "
        "sensible a la penetración. Los coeficientes del prototipo son didácticos y deben calibrarse con mediciones.",
    )
    add_no_indent(doc, "Palabras clave: FSPL, atenuación pasiva, radiofrecuencia, blindaje electromagnético, centros penitenciarios.", bold=True)
    add_heading(doc, "Abstract", 1)
    add_body(
        doc,
        "This project develops RF Shield Planner, a web simulator that estimates received power in a two-dimensional "
        "correctional-facility plan by combining free-space path loss and material attenuation. Controlled scenarios "
        "show that 850 MHz is harder to contain than 28 GHz. The prototype supports conceptual planning and requires "
        "field calibration before any real implementation.",
    )
    add_no_indent(doc, "Keywords: FSPL, passive attenuation, radio frequency, electromagnetic shielding, correctional facilities.", bold=True)

    # Índice estático
    doc.add_page_break()
    add_heading(doc, "Índice general", 1)
    toc_items = [
        ("Resumen", "3"),
        ("Introducción", "5"),
        ("Capítulo I Planteamiento del problema", "6"),
        ("Capítulo II Marco teórico", "8"),
        ("Capítulo III Marco metodológico", "11"),
        ("Capítulo IV Presentación y análisis de resultados", "14"),
        ("Conclusiones", "16"),
        ("Recomendaciones", "17"),
        ("Referencias", "18"),
        ("Anexos", "19"),
    ]
    for name, page in toc_items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.line_spacing = 1.5
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(Inches(6.1))
        run = paragraph.add_run(f"{name}\t{page}")
        set_font(run)
    add_heading(doc, "Índice de tablas", 1)
    for name in (
        "Tabla 1 Coeficientes efectivos del simulador",
        "Tabla 2 Variables de estudio",
        "Tabla 3 Escenarios de prueba",
        "Tabla 4 Resultados comparativos",
    ):
        add_no_indent(doc, name)
    add_heading(doc, "Índice de figuras", 1)
    add_no_indent(doc, "Figura 1 Porcentaje de área interior bajo -95 dBm")

    # Introducción
    doc.add_page_break()
    add_heading(doc, "Introducción", 1)
    add_body(
        doc,
        "Las redes móviles son infraestructura crítica para la actividad económica, la atención de emergencias y la "
        "comunicación cotidiana. Esa misma disponibilidad puede ser aprovechada de manera no autorizada desde centros "
        "penitenciarios. El problema exige una respuesta técnica que reduzca la conectividad dentro del objetivo sin "
        "interrumpir las comunicaciones legítimas del entorno.",
    )
    add_body(
        doc,
        "Los inhibidores activos introducen energía en bandas utilizadas por terceros y pueden producir desborde fuera "
        "del perímetro previsto. Por esa razón, este trabajo estudia la atenuación pasiva: el uso de materiales, masa, "
        "capas y continuidad constructiva para reducir la potencia recibida. La propuesta se materializa en una "
        "aplicación web que permite dibujar escenarios y comparar bandas móviles con un mismo balance de potencia.",
    )
    add_body(
        doc,
        "El documento se organiza en cuatro capítulos. El primero delimita el problema, los objetivos y la justificación. "
        "El segundo desarrolla los fundamentos de propagación, longitud de onda, FSPL, balance de potencia, materiales y "
        "marco legal. El tercero explica la metodología y la implementación del simulador. El cuarto presenta los "
        "resultados, la discusión y el análisis So What. Finalmente se formulan conclusiones, recomendaciones y anexos "
        "para reproducir la demostración.",
    )

    # Capítulo I
    add_chapter(doc, "I", "Planteamiento del problema")
    add_heading(doc, "1.1 Antecedentes", 2)
    add_body(
        doc,
        "La expansión de las redes celulares incrementó la disponibilidad de servicios de voz y datos en prácticamente "
        "todo entorno urbano. Sin embargo, los límites administrativos de un centro penitenciario no coinciden con los "
        "límites físicos de cobertura. Una estación base exterior puede mantener servicio en patios, celdas y áreas de "
        "visita cuando la pérdida de propagación y penetración no supera el margen del enlace.",
    )
    add_body(
        doc,
        "La respuesta tradicional basada en jamming busca elevar artificialmente el ruido o bloquear canales de control. "
        "Su principal dificultad es el confinamiento espacial. La radiación interferente puede atravesar el perímetro y "
        "afectar a residentes, comercios, operadores y servicios de emergencia. La alternativa estudiada consiste en "
        "reducir la potencia de señal mediante la propia envolvente del edificio.",
    )
    add_heading(doc, "1.2 Definición del problema", 2)
    add_body(
        doc,
        "No existe una relación intuitiva única entre frecuencia, distancia, material y grosor que permita decidir si un "
        "cerramiento llevará una señal celular por debajo de un umbral de comunicación. Sin una herramienta comparativa, "
        "es fácil asumir que cualquier muro produce el mismo efecto en 850 MHz, 1900 MHz, 3.5 GHz y 28 GHz.",
    )
    add_heading(doc, "1.3 Pregunta de investigación", 2)
    add_body(
        doc,
        "¿Cómo puede un simulador web conceptual demostrar el efecto combinado de FSPL, frecuencia, material y grosor "
        "sobre la potencia recibida dentro de un centro penitenciario, manteniendo como criterio un umbral de -95 dBm?",
    )
    add_heading(doc, "1.4 Objetivo general", 2)
    add_body(
        doc,
        "Diseñar e implementar un simulador web que represente escenarios penitenciarios bidimensionales y estime el "
        "confinamiento pasivo de señales celulares mediante FSPL y pérdidas efectivas por materiales constructivos.",
    )
    add_heading(doc, "1.5 Objetivos específicos", 2)
    for item in (
        "Implementar el cálculo de FSPL según distancia y frecuencia.",
        "Representar la relación inversa entre frecuencia y longitud de onda.",
        "Permitir la creación de planos con muros, puertas, rejas y zonas interiores.",
        "Calcular pérdidas acumuladas cuando el trayecto atraviesa uno o varios obstáculos.",
        "Comparar 850 MHz, 1900 MHz, 2100 MHz, 3.5 GHz y 28 GHz.",
        "Evaluar el porcentaje del área interior que queda bajo -95 dBm.",
        "Justificar la atenuación pasiva desde una perspectiva ética y legal.",
    ):
        add_bullet(doc, item)
    add_heading(doc, "1.6 Justificación", 2)
    add_body(
        doc,
        "El simulador facilita una comprensión cuantitativa del problema y convierte conceptos abstractos en decisiones "
        "observables. El usuario puede mantener fijo el plano y modificar una variable por vez, lo cual permite analizar "
        "la sensibilidad del resultado. Además, la separación entre frontend y backend favorece la claridad del código, "
        "la verificación de las fórmulas y un despliegue web posterior.",
    )
    add_heading(doc, "1.7 Alcances y delimitaciones", 2)
    add_body(
        doc,
        "El alcance incluye una aplicación web local, un editor 2D, cinco bandas preconfiguradas, cuatro tipos de "
        "obstáculos, mapa de potencia y resultados numéricos. Se excluyen propagación tridimensional, multitrayectoria, "
        "difracción rigurosa, medición de campo, costos constructivos y emisión activa. El prototipo es una herramienta "
        "académica de prefactibilidad.",
    )

    # Capítulo II
    add_chapter(doc, "II", "Marco teórico")
    add_heading(doc, "2.1 Espectro radioeléctrico y comunicaciones móviles", 2)
    add_body(
        doc,
        "El espectro radioeléctrico es un recurso limitado administrado por el Estado. Las redes móviles utilizan bandas "
        "con características de propagación distintas. Una frecuencia no determina por sí sola la cobertura, pero "
        "interviene en la pérdida de espacio libre, la longitud de onda, la interacción con materiales y el tamaño de "
        "aberturas capaces de permitir fuga electromagnética.",
    )
    add_heading(doc, "2.2 Longitud de onda", 2)
    add_body(doc, "La longitud de onda se calcula mediante la relación entre velocidad de propagación y frecuencia.")
    add_equation(doc, "λ = c / f", "Ecuación 1 Longitud de onda")
    add_body(
        doc,
        "Con una velocidad de la luz de 299 792 458 m/s, la longitud de onda es aproximadamente 35.27 cm a 850 MHz, "
        "15.78 cm a 1900 MHz, 8.57 cm a 3.5 GHz y 1.07 cm a 28 GHz. Una longitud menor incrementa la sensibilidad a "
        "detalles geométricos, discontinuidades y materiales del entorno.",
    )
    add_heading(doc, "2.3 Pérdida de espacio libre", 2)
    add_body(
        doc,
        "FSPL representa la reducción ideal de densidad de potencia entre transmisor y receptor sin obstáculos. Para "
        "distancia en kilómetros y frecuencia en megahercios se emplea la siguiente expresión:",
    )
    add_equation(
        doc,
        "FSPL (dB) = 32.44 + 20 log₁₀(dₖₘ) + 20 log₁₀(fₘₕz)",
        "Ecuación 2 Pérdida de espacio libre",
    )
    add_body(
        doc,
        "El término de frecuencia muestra que, a igual distancia y ganancias, 28 GHz presenta una pérdida de espacio "
        "libre mayor que 850 MHz. Esto no equivale a afirmar que toda banda alta tenga automáticamente menos cobertura, "
        "porque las redes pueden compensar con ganancia, densidad de estaciones, potencia y formación de haces.",
    )
    add_heading(doc, "2.4 Balance de potencia y relación señal a ruido", 2)
    add_equation(
        doc,
        "Pᵣ = Pₜ + Gₜ + Gᵣ - FSPL - ∑ Lₘₐₜₑᵣᵢₐₗ",
        "Ecuación 3 Balance de potencia recibido",
    )
    add_body(
        doc,
        "La potencia recibida se compara con -95 dBm como umbral conceptual. El simulador mantiene un piso de ruido de "
        "-110 dBm, por lo que la relación señal a ruido se obtiene restando el nivel de ruido a la potencia recibida. La "
        "atenuación pasiva reduce primero la señal; si el ruido permanece fijo, también disminuye la SNR.",
    )
    add_heading(doc, "2.5 Atenuación por materiales", 2)
    add_equation(
        doc,
        "Lₘₐₜₑᵣᵢₐₗ = α(f, material) · t",
        "Ecuación 4 Pérdida efectiva por material",
    )
    add_body(
        doc,
        "El coeficiente efectivo alpha depende de la frecuencia y del material, mientras t representa el grosor. Esta "
        "aproximación facilita la demostración, pero no es una constante física universal. La UIT-R P.2040 describe "
        "propiedades eléctricas, interfaces, losas y estructuras multicapa. NISTIR 6055 muestra mediante mediciones que "
        "la pérdida cambia con frecuencia, espesor, humedad, mezcla y refuerzo.",
    )
    doc.add_paragraph("Tabla 1 Coeficientes efectivos utilizados por el prototipo", style="Table Caption")
    add_table(
        doc,
        ["Material", "850 MHz", "1.9 GHz", "2.1 GHz", "3.5 GHz", "28 GHz"],
        coefficient_rows(),
        [Inches(1.55), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9), Inches(0.9)],
    )
    add_body(
        doc,
        "Los valores se expresan en dB/m y constituyen parámetros didácticos de calibración. Una recomendación "
        "constructiva debe reemplazarlos con ensayos o bibliografía aplicable al material real, incluyendo humedad, "
        "armadura, juntas, ventanas, puertas, ductos y ángulo de incidencia.",
        "Los valores",
    )
    add_heading(doc, "2.6 Blindaje electromagnético y jaula de Faraday", 2)
    add_body(
        doc,
        "Un blindaje conductor redistribuye cargas y corrientes de manera que reduce el campo electromagnético dentro "
        "del volumen protegido. Su eficacia depende de conductividad, continuidad, espesor, uniones y tamaño de las "
        "aberturas respecto de la longitud de onda. Una puerta mal conectada, una rejilla grande o un ducto sin tratamiento "
        "pueden dominar la fuga aun cuando los paneles tengan alta pérdida de inserción.",
    )
    add_heading(doc, "2.7 Marco legal y ético", 2)
    add_body(
        doc,
        "La Ley General de Telecomunicaciones de Guatemala establece un marco para el aprovechamiento del espectro, la "
        "protección de usuarios y empresas proveedoras y el uso racional del recurso. La SIT dispone procedimientos para "
        "la atención de interferencias perjudiciales. Desde esta perspectiva, una solución que emite ruido fuera del "
        "objetivo introduce riesgo para terceros y debe evitarse en un proyecto académico.",
    )
    add_body(
        doc,
        "La atenuación pasiva no ocupa una banda mediante una nueva transmisión. Su propósito es actuar sobre la "
        "envolvente física para reducir conectividad dentro del recinto y conservar las comunicaciones autorizadas del "
        "exterior. Una implementación real deberá preservar canales de emergencia y coordinarse con la SIT, operadores y "
        "autoridades penitenciarias.",
    )

    # Capítulo III
    add_chapter(doc, "III", "Marco metodológico")
    add_heading(doc, "3.1 Enfoque y tipo de estudio", 2)
    add_body(
        doc,
        "La investigación adopta un enfoque cuantitativo aplicado y un diseño experimental por simulación. No se manipula "
        "espectro real. Los escenarios se ejecutan con parámetros controlados y se compara la respuesta numérica del "
        "modelo al modificar frecuencia, material o grosor.",
    )
    add_heading(doc, "3.2 Variables de estudio", 2)
    doc.add_paragraph("Tabla 2 Variables de estudio", style="Table Caption")
    add_table(
        doc,
        ["Variable", "Tipo", "Unidad o categoría", "Función"],
        [
            ["Frecuencia", "Independiente", "MHz", "Modifica FSPL, longitud de onda y alpha."],
            ["Distancia", "Independiente", "m", "Modifica FSPL."],
            ["Potencia transmitida", "Independiente", "dBm", "Punto de partida del balance."],
            ["Material", "Independiente", "Categoría", "Selecciona coeficiente efectivo."],
            ["Grosor", "Independiente", "m", "Escala la pérdida del obstáculo."],
            ["Potencia recibida", "Dependiente", "dBm", "Resultado por punto del plano."],
            ["Área confinada", "Dependiente", "%", "Puntos interiores bajo -95 dBm."],
        ],
        [Inches(1.35), Inches(1.15), Inches(1.35), Inches(3.0)],
    )
    add_heading(doc, "3.3 Arquitectura del sistema", 2)
    add_body(
        doc,
        "El frontend y el backend se desarrollaron como aplicaciones separadas. El frontend contiene el editor Canvas, "
        "medidas, edición de propiedades, almacenamiento local, comparación de bandas y visualización del balance. El "
        "backend centraliza materiales, validación y cálculo. La comunicación utiliza "
        "JSON a través del endpoint POST /api/simulate. Esta separación permite desplegar la interfaz como contenido "
        "estático y la API en un servicio independiente.",
    )
    add_table(
        doc,
        ["Componente", "Responsabilidad", "Archivo"],
        [
            ["Editor 2D", "Dibujar, medir, mover y editar elementos.", "frontend/src/editor.js"],
            ["Interfaz", "Gestionar controles, comparación y guardado.", "frontend/src/main.js"],
            ["Cliente API", "Enviar y recibir JSON.", "frontend/src/api.js"],
            ["Servidor", "Exponer endpoints y validar solicitudes.", "backend/src/server.js"],
            ["Modelo RF", "Calcular FSPL, intersecciones y potencia.", "backend/src/rf-model.js"],
            ["Materiales", "Centralizar parámetros del modelo.", "backend/src/materials.js"],
        ],
        [Inches(1.35), Inches(3.35), Inches(2.2)],
    )
    add_heading(doc, "3.4 Algoritmo de simulación", 2)
    for step in (
        "Dividir el lienzo en una cuadrícula de puntos de evaluación.",
        "Calcular la distancia desde la antena hasta cada punto.",
        "Aplicar FSPL para la frecuencia seleccionada.",
        "Detectar los segmentos del plano cruzados por el trayecto.",
        "Sumar la pérdida efectiva de cada obstáculo según material y grosor.",
        "Obtener potencia recibida y SNR.",
        "Clasificar el punto respecto de -95 dBm.",
        "Calcular el porcentaje de puntos interiores bajo el umbral.",
        "Presentar el balance promedio y comparar 850 MHz con 28 GHz.",
    ):
        add_bullet(doc, step)
    add_heading(doc, "3.5 Escenarios de prueba", 2)
    doc.add_paragraph("Tabla 3 Escenarios de prueba", style="Table Caption")
    add_table(
        doc,
        ["Escenario", "Frecuencia", "Construcción", "Propósito"],
        [
            ["A", "850 MHz", "Concreto de 0.25 m", "Demostrar dificultad de confinamiento."],
            ["B", "1900 MHz", "Concreto de 0.25 m", "Establecer línea base intermedia."],
            ["C", "28 GHz", "Mismo plano de concreto", "Comparar mmWave sin cambiar geometría."],
            ["D", "1900 MHz", "Concreto más barrera metálica", "Evaluar refuerzo de frontera crítica."],
        ],
        [Inches(0.75), Inches(1.15), Inches(2.35), Inches(2.75)],
    )
    add_heading(doc, "3.6 Procedimiento", 2)
    add_body(
        doc,
        "Se carga el plano de ejemplo, se fija la potencia transmitida en 43 dBm, la distancia base en un kilómetro y el "
        "umbral en -95 dBm. Se ejecutan 850 MHz, 1900 MHz y 28 GHz sin modificar la geometría. Después se restaura el "
        "caso de 1900 MHz, se añade una barrera metálica interior de 0.25 m en el modelo y se repite el cálculo. Para cada "
        "ejecución se registran potencia promedio, potencia mínima, longitud de onda, FSPL, pérdida por obstáculos y "
        "porcentaje confinado. El comparador automatiza la ejecución de 850 MHz y 28 GHz con el mismo plano.",
    )
    add_heading(doc, "3.7 Validación", 2)
    add_body(
        doc,
        "La validación lógica incluye pruebas unitarias para FSPL a 1 km y 1000 MHz, relación inversa entre frecuencia y "
        "longitud de onda, crecimiento de pérdida con grosor, intersección de segmentos, tamaño de la cuadrícula y mayor "
        "atenuación efectiva del concreto a 28 GHz que a 850 MHz. La compilación del frontend y la respuesta del endpoint "
        "también se verifican antes de la demostración.",
    )

    # Capítulo IV
    add_chapter(doc, "IV", "Presentación y análisis de resultados")
    add_heading(doc, "4.1 Resultados cuantitativos", 2)
    doc.add_paragraph("Tabla 4 Resultados comparativos", style="Table Caption")
    add_table(
        doc,
        ["Escenario", "Promedio", "Mínimo", "Área bajo -95 dBm", "Longitud de onda"],
        [
            ["850 MHz con concreto", "-56.7 dBm", "-61.1 dBm", "0 %", "35.27 cm"],
            ["1900 MHz con concreto", "-67.6 dBm", "-76.0 dBm aprox.", "0 %", "15.78 cm"],
            ["28 GHz con concreto", "-124.7 dBm", "-141.9 dBm", "100 %", "1.07 cm"],
            ["1900 MHz con blindaje", "-114.2 dBm", "-122.7 dBm", "94.4 %", "15.78 cm"],
        ],
        [Inches(1.65), Inches(1.05), Inches(1.15), Inches(1.4), Inches(1.35)],
    )
    add_figure(doc, chart, "Figura 1 Porcentaje de área interior bajo -95 dBm", Inches(6.0))
    add_heading(doc, "4.2 Escenario de 850 MHz", 2)
    add_body(
        doc,
        "El resultado de 0 % bajo el umbral indica que el cerramiento convencional no produce la pérdida necesaria. La "
        "potencia promedio de -56.7 dBm conserva un margen de 38.3 dB respecto de -95 dBm. En consecuencia, una solución "
        "pasiva debe acumular varias capas, mayor grosor o un sistema conductor continuo, además de controlar aberturas.",
    )
    add_heading(doc, "4.3 Escenario de 28 GHz", 2)
    add_body(
        doc,
        "Al conservar el mismo plano y cambiar únicamente la frecuencia, el modelo lleva el 100 % del área interior bajo "
        "el umbral. La potencia promedio cae a -124.7 dBm. La diferencia combina el aumento de FSPL con una mayor "
        "atenuación efectiva asignada a los materiales. El resultado demuestra el comportamiento esperado para ondas "
        "milimétricas, aunque no autoriza generalizaciones sin medición.",
    )
    add_heading(doc, "4.4 Escenario de 1900 MHz con refuerzo", 2)
    add_body(
        doc,
        "El caso inicial de 1900 MHz mantiene toda el área sobre el umbral. Al añadir una barrera metálica conceptual, el "
        "porcentaje confinado aumenta a 94.4 % y la potencia promedio disminuye a -114.2 dBm. El mapa permite localizar el "
        "5.6 % restante y estudiar si corresponde a accesos, geometría o trayectos que no atraviesan la nueva barrera.",
    )
    add_heading(doc, "4.5 Análisis So What", 2)
    add_body(
        doc,
        "El hallazgo operativo es que una estrategia uniforme no funciona para todas las bandas. A 850 MHz, un muro "
        "convencional deja un margen amplio de comunicación; por ello se requiere un diseño arquitectónico pesado, "
        "multicapa y continuo. A 28 GHz, el enlace dispone de menos margen para penetrar, de modo que los materiales "
        "básicos producen confinamiento casi perfecto dentro del escenario modelado.",
    )
    add_body(
        doc,
        "La decisión de ingeniería no consiste únicamente en elegir el material con el mayor coeficiente. También debe "
        "considerar continuidad, accesos, ventilación, estructura, costos, seguridad contra incendios y canales autorizados. "
        "El simulador sirve para identificar hipótesis y priorizar mediciones, no para reemplazar el diseño detallado.",
    )
    add_heading(doc, "4.6 Limitaciones de los resultados", 2)
    for item in (
        "El plano no representa techos, sótanos ni pisos superiores.",
        "No se modelan reflexiones, difracción detallada, polarización ni desvanecimiento.",
        "Los coeficientes son parámetros didácticos y no mediciones de un material específico.",
        "El umbral de -95 dBm es conceptual y la sensibilidad real depende del equipo y la red.",
        "No se modelan simultáneamente enlace ascendente y descendente.",
    ):
        add_bullet(doc, item)

    # Conclusiones
    doc.add_page_break()
    add_heading(doc, "Conclusiones", 1)
    conclusions = [
        "El prototipo demuestra que la potencia recibida puede estimarse mediante un balance que combina potencia transmitida, FSPL y pérdidas acumuladas por obstáculos.",
        "El escenario de 850 MHz permanece por encima de -95 dBm con concreto de 0.25 m, lo cual respalda la necesidad de una solución arquitectónica pesada para frecuencias bajas.",
        "El mismo plano alcanza 100 % de confinamiento a 28 GHz dentro del modelo, debido al aumento de FSPL y a la mayor pérdida efectiva de penetración.",
        "Una barrera metálica conceptual incrementa el confinamiento de 1900 MHz hasta aproximadamente 94.4 %, demostrando el valor de reforzar fronteras críticas.",
        "La atenuación pasiva es éticamente preferible a la interferencia activa porque no introduce una nueva emisión destinada a degradar comunicaciones de terceros.",
        "Los coeficientes didácticos no constituyen especificaciones constructivas; antes de cualquier aplicación física deben calibrarse con mediciones y fuentes aplicables al sitio.",
    ]
    for index, item in enumerate(conclusions, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(f"{index}. {item}")
        set_font(run)

    doc.add_page_break()
    add_heading(doc, "Recomendaciones", 1)
    recommendations = [
        "Medir potencia exterior e interior en las bandas de interés antes de definir materiales o espesores.",
        "Sustituir los coeficientes didácticos por valores obtenidos de muestras reales, considerando humedad y refuerzo.",
        "Modelar puertas, ventanas, ductos, juntas y otras discontinuidades como elementos independientes.",
        "Extender el simulador a tres dimensiones e incorporar multitrayectoria para una segunda fase.",
        "Verificar que cualquier solución conserve comunicaciones autorizadas y servicios de emergencia.",
        "Coordinar las pruebas físicas con la SIT, operadores y autoridades penitenciarias competentes.",
        "Utilizar el simulador como herramienta de prefactibilidad y enseñanza, no como certificación de blindaje.",
    ]
    for index, item in enumerate(recommendations, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.65)
        paragraph.paragraph_format.first_line_indent = Cm(-0.65)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(f"{index}. {item}")
        set_font(run)

    # Referencias
    doc.add_page_break()
    add_heading(doc, "Referencias", 1)
    references = [
        "Congreso de la República de Guatemala. (1996). Decreto 94-96, Ley General de Telecomunicaciones. https://www.congreso.gob.gt/detalle_pdf/decretos/897",
        "Du, K., Ozdemir, O., Erden, F. y Guvenc, I. (2021). Sub-Terahertz and mmWave penetration loss measurements for indoor environments. IEEE ICC Workshops. https://doi.org/10.1109/ICCWorkshops50388.2021.9473898",
        "Stone, W. C. (1997). Electromagnetic signal attenuation in construction materials. NISTIR 6055. National Institute of Standards and Technology. https://doi.org/10.6028/NIST.IR.6055",
        "Superintendencia de Telecomunicaciones de Guatemala. (s. f.). Supervisión de espectro. https://sit.gob.gt/supervision-de-espectro.html",
        "Superintendencia de Telecomunicaciones de Guatemala. (s. f.). Ley General de Telecomunicaciones. https://sit.gob.gt/files/LEY-GENERAL-DE-TELECOMUNICACIONES.pdf",
        "Unión Internacional de Telecomunicaciones. (2025). Recomendación UIT-R P.2040-4: Effects of building materials and structures on radiowave propagation in the range of 1 MHz to 450 GHz. https://www.itu.int/rec/R-REC-P.2040",
    ]
    for reference in references:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(1.25)
        paragraph.paragraph_format.first_line_indent = Cm(-1.25)
        paragraph.paragraph_format.line_spacing = 1.5
        run = paragraph.add_run(reference)
        set_font(run)

    # Anexos
    doc.add_page_break()
    add_heading(doc, "Anexos", 1)
    add_heading(doc, "Anexo A Manual breve del simulador", 2)
    manual = [
        "Iniciar el backend con npm start dentro de la carpeta backend.",
        "Iniciar el frontend con npm run dev dentro de la carpeta frontend.",
        "Abrir http://127.0.0.1:5173 en el navegador.",
        "Seleccionar Piso o zona y dibujar el área que se evaluará.",
        "Seleccionar material y grosor, elegir Muro y dibujar el perímetro.",
        "Añadir puertas y rejas sobre los muros; el editor abre el tramo correspondiente.",
        "Usar Seleccionar para mover elementos, cambiar propiedades o eliminarlos.",
        "Colocar la antena fuera del edificio.",
        "Definir frecuencia, potencia, distancia y umbral.",
        "Ejecutar Simular cobertura y explicar el balance mostrado.",
        "Ejecutar Comparar 850 MHz vs 28 GHz y registrar la diferencia.",
        "Guardar el plano en el navegador cuando se desee reutilizarlo.",
    ]
    for index, item in enumerate(manual, 1):
        paragraph = add_no_indent(doc, f"{index}. {item}")
        paragraph.paragraph_format.line_spacing = 1.15
    add_heading(doc, "Anexo B Demostración sugerida", 2)
    add_body(
        doc,
        "La demostración en video debe mantener el plano, potencia, distancia y umbral constantes. Primero se ejecuta "
        "850 MHz, luego 28 GHz. Finalmente se restaura 1900 MHz, se registra la línea base y se añade la barrera "
        "metálica. El video debe mostrar el panel Balance calculado, la comparación automática y las funciones "
        "calculateFspl, calculateObstacleLoss y calculatePoint para relacionar la interfaz con el código.",
    )
    add_heading(doc, "Anexo C Estructura del repositorio", 2)
    structure = [
        ("frontend", "Interfaz, editor Canvas y cliente de API."),
        ("backend", "Servidor Express, materiales, modelo RF y pruebas."),
        ("entregables", "Guion de video y lista de control."),
        ("output/docx", "Documento de tesis editable."),
        ("output/pdf", "Informe y versión PDF de la tesis."),
        ("scripts", "Generadores reproducibles de documentos."),
    ]
    add_table(
        doc,
        ["Carpeta", "Contenido"],
        structure,
        [Inches(1.7), Inches(4.8)],
        trailing_space=False,
    )

    doc.core_properties.title = "Diseño de un simulador web para el análisis conceptual de atenuación pasiva"
    doc.core_properties.subject = "Proyecto de tesis en telecomunicaciones"
    doc.core_properties.author = "Estudiante"
    doc.core_properties.keywords = "FSPL, atenuación pasiva, radiofrecuencia, telecomunicaciones"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
